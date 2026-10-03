"""Analysis exports must never mutate the Power BI Project being inspected."""
import json
import os
import subprocess
import pytest
from openpyxl import load_workbook

from semantic_model_cleaner import cli


@pytest.fixture
def project(tmp_path):
    root = tmp_path / "Power BI Project"
    model = root / "Sales.SemanticModel"
    report = root / "Overview.Report"
    excluded_model = root / "Excluded.SemanticModel"
    excluded_report = root / "Excluded.Report"
    for artifact in (model, excluded_model):
        (artifact / "definition/tables").mkdir(parents=True)
        (artifact / "definition/tables/Sales.tmdl").write_text(
            "table Sales\n\tmeasure Revenue = 42\n", encoding="utf-8"
        )
    for artifact in (report, excluded_report):
        (artifact / "definition").mkdir(parents=True)
        (artifact / "definition/report.json").write_text("{}", encoding="utf-8")
        (artifact / "definition.pbir").write_text(json.dumps({
            "datasetReference": {"byPath": {"path": "../Sales.SemanticModel"}}
        }), encoding="utf-8")
    return root, model, report, excluded_model, excluded_report


def snapshot(root):
    # Include directory entries as well as file bytes, so even a recovery folder
    # or new export created by a rejected operation makes this check fail.
    return {str(path.relative_to(root)): path.read_bytes() if path.is_file() else None
            for path in root.rglob("*")}


def make_directory_alias(alias, target):
    if os.name == "nt":
        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command",
             "$ErrorActionPreference = 'Stop'; New-Item -ItemType Junction "
             "-Path $env:SMC_TEST_ALIAS -Target $env:SMC_TEST_TARGET | Out-Null"],
            env={**os.environ, "SMC_TEST_ALIAS": str(alias), "SMC_TEST_TARGET": str(target)},
            capture_output=True, text=True,
        )
        assert result.returncode == 0, result.stderr
    else:
        alias.symlink_to(target, target_is_directory=True)


def export(project, output, format="json", extra=()):
    args = [str(project[0]), "--model", "Sales", "--report", "Overview", "--format", format]
    if output is not None:
        args += ["--output", str(output)]
    cli.main([*args, *extra])


@pytest.mark.parametrize("format", ["full", "unused", "json", "xlsx"])
@pytest.mark.parametrize("artifact_index", [1, 2, 3, 4])
@pytest.mark.parametrize("existing", [True, False])
def test_export_rejects_selected_and_excluded_artifacts(project, capsys, format, artifact_index, existing):
    artifact = project[artifact_index]
    relative = "definition/tables/Sales.tmdl" if artifact_index in (1, 3) else "definition/report.json"
    output = artifact / (relative if existing else "new-export")
    before = snapshot(project[0])
    with pytest.raises(SystemExit) as error:
        export(project, output, format)
    assert error.value.code == 2
    assert "outside" in capsys.readouterr().err
    assert snapshot(project[0]) == before


@pytest.mark.parametrize("artifact_index", [1, 2])
def test_default_excel_output_rejects_working_directory_in_artifact(project, monkeypatch, artifact_index):
    monkeypatch.chdir(project[artifact_index] / "definition")
    before = snapshot(project[0])
    with pytest.raises(SystemExit) as error:
        export(project, None, "xlsx")
    assert error.value.code == 2
    assert snapshot(project[0]) == before


@pytest.mark.parametrize("format", ["full", "unused", "json", "xlsx"])
def test_external_export_and_overwrite_succeed(project, tmp_path, format):
    destination = tmp_path / "Zażółć exports"
    destination.mkdir()
    output = destination / f"analysis.{format}"
    output.write_text("old export", encoding="utf-8")
    before = snapshot(project[0])
    export(project, output, format)
    assert snapshot(project[0]) == before
    if format == "json":
        assert json.loads(output.read_text(encoding="utf-8"))["items"]
    elif format == "xlsx":
        workbook = load_workbook(output)
        assert "Summary" in workbook.sheetnames
        workbook.close()
    else:
        assert "Revenue" in output.read_text(encoding="utf-8")


@pytest.mark.parametrize("format", ["full", "unused", "json", "xlsx"])
def test_external_hard_link_cannot_overwrite_source(project, tmp_path, capsys, format):
    source = project[1] / "definition/tables/Sales.tmdl"
    output = tmp_path / "hard-link"
    try:
        output.hardlink_to(source)
    except OSError as exc:
        pytest.skip(f"Hard links unavailable: {exc}")
    before = snapshot(tmp_path)
    with pytest.raises(SystemExit) as error:
        export(project, output, format)
    assert error.value.code == 2
    assert "hard link" in capsys.readouterr().err
    assert snapshot(tmp_path) == before


@pytest.mark.parametrize("kind", ["file_symlink", "directory_symlink", "junction"])
def test_alias_cannot_overwrite_artifact(project, tmp_path, kind):
    alias = tmp_path / "alias"
    if kind == "junction":
        if os.name != "nt":
            pytest.skip("Windows junction test")
        make_directory_alias(alias, project[1])
        output = alias / "definition/tables/Sales.tmdl"
    else:
        source = project[1] if kind == "directory_symlink" else project[1] / "definition/tables/Sales.tmdl"
        try:
            alias.symlink_to(source, target_is_directory=kind == "directory_symlink")
        except OSError as exc:
            pytest.skip(f"Symlinks unavailable: {exc}")
        output = alias / "definition/tables/Sales.tmdl" if kind == "directory_symlink" else alias
    before = snapshot(project[0])
    with pytest.raises(SystemExit) as error:
        export(project, output)
    assert error.value.code == 2
    assert snapshot(project[0]) == before


def test_discovered_excluded_alias_protects_unnamed_target(project, tmp_path):
    target = tmp_path / "metadata storage"
    target.mkdir()
    output = target / "metadata.json"
    output.write_text("original", encoding="utf-8")
    make_directory_alias(project[0] / "Another.Report", target)
    before = snapshot(tmp_path)
    with pytest.raises(SystemExit) as error:
        export(project, output)
    assert error.value.code == 2
    assert snapshot(tmp_path) == before


def test_export_with_parent_traversal_to_external_directory_succeeds(project):
    output = project[1] / ".." / "analysis.json"
    export(project, output)
    assert json.loads(output.read_text(encoding="utf-8"))["items"]


def test_default_excel_external_directory_succeeds(project, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    before = snapshot(project[0])
    export(project, None, "xlsx")
    assert (tmp_path / "Sales_usage_analysis.xlsx").is_file()
    assert snapshot(project[0]) == before


def test_explicit_artifact_roots_outside_workspace_remain_protected(project, tmp_path):
    workspace = tmp_path / "separate workspace"
    workspace.mkdir()
    before = snapshot(project[0])
    with pytest.raises(SystemExit) as error:
        cli.main([str(workspace), "--models-path", str(project[1]), "--reports-path", str(project[2]),
                  "--format", "json", "--output", str(project[1] / "definition/tables/Sales.tmdl")])
    assert error.value.code == 2
    assert snapshot(project[0]) == before
