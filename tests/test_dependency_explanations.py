"""Declaration-based feature detection, dependency explanations and analysis limitations.

Covers GitHub issues #82 (detect features from TMDL declarations), #83 (explain
calculation-group and perspective dependencies) and #84 (separate analysis
limitations from Report Health). The synthetic fixture has a calculation group
with ten retained parent-table consumers, a table named `KPI Selector` that
declares no KPI, a genuine KPI, and perspective membership.
"""
import json
from pathlib import Path

import pytest

from semantic_model_cleaner import analyzer, ci, webapp
from semantic_model_cleaner.cleanup_policy import evaluate_deletion_policy
from semantic_model_cleaner.tmdl_declarations import (
    extract_feature_expressions,
    extract_perspective_members,
    scan_tmdl_declarations,
)

FIXTURE = Path(__file__).parent / "fixtures" / "calculation_group_dependencies"
MODEL = FIXTURE / "Models" / "Synthetic Dependencies.SemanticModel"
REPORT = FIXTURE / "Reports" / "Executive.Report"
CONSUMERS = [f"Sales[Revenue Ignoring TI {n:02d}]" for n in range(1, 11)]


@pytest.fixture(scope="module")
def results():
    return analyzer.analyze(FIXTURE, model_paths=[MODEL], report_paths=[REPORT])


@pytest.fixture(scope="module")
def payload(results):
    return webapp._serialize_results(results, model_paths=[str(MODEL)])


def row(results, table, name):
    return next(r for r in results["items"] if r["item"].table == table and r["item"].name == name)


def item(payload, table, name):
    return next(i for i in payload["items"] if i["table"] == table and i["name"] == name)


# ── #82: detection from declarations, not words ──────────────────────────────

@pytest.mark.parametrize("text", [
    "table 'KPI Selector'\n\tcolumn Name\n\t\tdataType: string\n",
    "table T\n\t/// kpi description\n\tmeasure M = 1\n\t\tdescription: targetExpression mentioned in prose\n",
    "table T\n\t// kpi comment\n\tmeasure M = \"kpi\"\n",
    "table T\n\tmeasure M =\n\t\t\tVAR kpi = 1\n\t\t\tRETURN kpi\n",
    "table T\n\tmeasure M = 1\n\t\tdisplayFolder: KPI\\calculationGroup\n",
    "table T\n\t/* kpi\n\tcalculationGroup */\n\tmeasure M = 1\n",
    "table T\n\tpartition P = m\n\t\tsource =\n\t\t\tlet kpi = 1, calculationItem = 2 in kpi\n",
])
def test_feature_words_outside_declarations_detect_nothing(text):
    features, tables = extract_feature_expressions(text)
    assert features == []
    assert tables == {}


def test_genuine_declarations_report_owner_and_location():
    text = (
        "table Targets\n\tmeasure Target = 1000\n\t\tkpi\n\t\t\ttargetExpression = 'Sales'[Goal]\n"
        "\t\t\tstatusExpression = ```\n\t\t\t\t\tVAR x = [Target]\n\t\t\t\t\tRETURN x\n\t\t\t\t```\n"
        "\t\tdetailRowsDefinition = SELECTCOLUMNS(Sales, \"x\", Sales[Amount])\n"
        "table 'Time Intelligence'\n\tcalculationGroup\n\t\tcalculationItem YTD = SELECTEDMEASURE()\n"
        "\t\t\tformatStringDefinition = SELECTEDMEASUREFORMATSTRING()\n"
    )
    features, tables = extract_feature_expressions(text)
    assert [(f.area, f.feature, f.owner, f.line) for f in features] == [
        ("KPI expressions", "KPI target expression", "measure 'Targets'[Target]", 4),
        ("KPI expressions", "KPI status expression", "measure 'Targets'[Target]", 5),
        ("Detail rows", "measure detail rows expression", "measure 'Targets'[Target]", 9),
        ("Calculation Groups", "calculation item expression", "calculation item 'YTD' in 'Time Intelligence'", 12),
        ("Format string definitions", "calculation item format-string expression",
         "calculation item 'YTD' in 'Time Intelligence'", 13),
    ]
    assert features[1].expression == "VAR x = [Target]\nRETURN x"
    assert tables == {"Time Intelligence": {"calculation_items": ["YTD"], "line": 11}}


def test_scanner_keeps_expression_continuations_out_of_the_tree():
    declarations = scan_tmdl_declarations(
        "table T\n\tmeasure M =\n\t\t\tVAR kpi = 1\n\t\t\tRETURN kpi\n\t\tformatString: 0\n")
    assert [(d.keyword, d.depth) for d in declarations] == [("table", 0), ("measure", 1), ("formatString", 2)]
    assert declarations[1].expression == "VAR kpi = 1\nRETURN kpi"


def _parse_items(tmp_path, text):
    tables = tmp_path / "M.SemanticModel" / "definition" / "tables"
    tables.mkdir(parents=True)
    (tables / "T.tmdl").write_text(text, encoding="utf-8")
    return {parsed.name: parsed for parsed in analyzer.parse_model_items(tmp_path / "M.SemanticModel")}


def test_kpi_child_object_is_not_folded_into_the_measure_dax_body():
    target = next(i for i in analyzer.parse_model_items(MODEL) if i.table == "Targets" and i.name == "Target")
    assert target.dax_body == "1000"


def test_item_dax_body_stops_at_nested_child_objects_but_keeps_continuations(tmp_path):
    items = _parse_items(tmp_path, (
        "table T\n"
        "\t/// Multi-line\n\t/// description\n"
        "\tmeasure Multi =\n\t\t\tVAR x = [A]\n\t\t\tRETURN x\n"
        "\t\tformatString: 0\n"
        "\t\tdisplayFolder: KPIs\n"
        "\t\tkpi\n\t\t\ttargetExpression = [Goal]\n\t\t\tstatusExpression =\n\t\t\t\t\t[Status]\n"
        "\t\tchangedProperty = Name\n"
        "\t\textendedProperty Ext =\n\t\t\t\t{ \"ref\": \"[Ext]\" }\n"
        "\t\tannotation Note = [Annotated]\n\n"
        "\tmeasure Fenced = ```\n\t\t\tVAR y = [B]\n\n\t\t\tRETURN y\n\t\t\t```\n"
        "\t\tkpi\n\t\t\ttrendExpression = [Trend]\n"
        "\t\tformatStringDefinition = [FormatRef]\n"
        "\t\tisHidden\n\n"
        "\tcolumn Calc = [C] + 1\n\t\tdataType: int64\n\t\tchangedProperty = DataType\n"
        "\t\tannotation Note =\n\t\t\t\t[ColumnAnnotation]\n"
        "\t\tisHidden\n\t\tisKey\n\t\tsortByColumn: 'Sort Key'\n\n"
        "\tcolumn Plain\n\t\tdataType: string\n\t\tsourceColumn: Plain\n\t\tisNameInferred\n"
    ))
    multi = items["Multi"]
    assert (multi.dax_body, multi.description, multi.format_string, multi.display_folder) == (
        "VAR x = [A]\nRETURN x", "Multi-line\ndescription", "0", "KPIs")
    fenced = items["Fenced"]
    assert (fenced.dax_body, fenced.is_hidden) == ("VAR y = [B]\nRETURN y\n[FormatRef]", True)
    calc = items["Calc"]
    assert (calc.item_type, calc.dax_body, calc.data_type, calc.is_hidden, calc.is_key, calc.sort_by_column) == (
        "Calculated Column", "[C] + 1", "int64", True, True, "Sort Key")
    plain = items["Plain"]
    assert (plain.item_type, plain.dax_body, plain.source_column, plain.is_inferred) == (
        "Column", "", "Plain", True)


# ── #96: calculated-column kind follows the structural expression ──────────

@pytest.mark.parametrize("declaration", [
    "\tcolumn Calc = [Base] + 1\n",
    "\tcolumn Calc =\n\t\t\t[Base] + 1\n",
    "\tcolumn Calc = ```\n\t\t\t[Base] + 1\n\t\t\t```\n",
    "\tcolumn Calc =\n\t\t\t```\n\t\t\t[Base] + 1\n\t\t\t```\n",
    "\tcolumn Calc\n\t\texpression =\n\t\t\t[Base] + 1\n",
], ids=["inline", "multiline", "inline-fence", "multiline-fence", "expression-property"])
def test_calculated_column_kind_does_not_depend_on_same_line_dax(tmp_path, declaration):
    items = _parse_items(tmp_path, (
        "table T\n"
        "\tcolumn Base\n\t\tdataType: int64\n\t\tsourceColumn: Base\n"
        "\t\tannotation Note = [NotDax]\n"
        + declaration + "\t\tdataType: int64\n\t\tdisplayFolder: Calculations\n"
        "\t\tannotation Note = [NotDax]\n"
    ))
    assert (items["Calc"].item_type, items["Calc"].dax_body, items["Calc"].data_type,
            items["Calc"].display_folder) == (
        "Calculated Column", "[Base] + 1", "int64", "Calculations")
    assert (items["Base"].item_type, items["Base"].dax_body, items["Base"].source_column) == (
        "Column", "", "Base")


MULTILINE_COLUMNS = (
    "table T\n"
    "\tcolumn Base\n\t\tdataType: int64\n\t\tsourceColumn: Base\n"
    "\tcolumn Multi =\n\t\t\tVAR value = [Base] + 1\n\t\t\tRETURN value\n"
    "\t\tdataType: int64\n"
    "\tcolumn Fenced =\n\t\t\t```\n\t\t\t[Base] + 2\n\t\t\t```\n"
    "\t\tdataType: int64\n"
    "\tcolumn Inline = [Base] + 3\n\t\tdataType: int64\n"
)


def test_multiline_calculated_column_kind_reaches_browser_json_and_xlsx(tmp_path):
    from openpyxl import load_workbook

    model, report = synthetic_project(tmp_path, {"T": MULTILINE_COLUMNS})
    analysis = analyzer.analyze(tmp_path, model_paths=[model], report_paths=[report])
    expected_types = {"Base": "Column", **dict.fromkeys(
        ["Multi", "Fenced", "Inline"], "Calculated Column")}
    assert {r["item"].name: r["item"].item_type for r in analysis["items"]} == expected_types
    assert analysis["summary"]["total_calc_columns"] == 3
    browser = webapp._serialize_results(analysis, model_paths=[str(model)])
    assert {i["name"]: i["type"] for i in browser["items"]} == expected_types
    assert item(browser, "T", "Multi")["daxExpression"] == "VAR value = [Base] + 1\nRETURN value"
    assert item(browser, "T", "Fenced")["daxExpression"] == "[Base] + 2"
    exported = json.loads(analyzer.format_json_output(analysis))
    assert {i["name"]: i["type"] for i in exported["items"]} == expected_types
    output = tmp_path / "analysis.xlsx"
    analyzer.format_xlsx(analysis, str(output), announce=False)
    workbook = load_workbook(output, read_only=True)
    try:
        rows = workbook["Details"].iter_rows(values_only=True)
        headers = next(rows)
        details = [dict(zip(headers, values)) for values in rows]
        assert {r["Name"]: r["Type"] for r in details} == expected_types
    finally:
        workbook.close()


def test_multiline_calculated_columns_keep_cleanup_dependency_guards(tmp_path):
    model, report = synthetic_project(tmp_path, {"T": MULTILINE_COLUMNS})
    base = {"action": "delete", "table": "T", "name": "Base", "item_type": "Column"}
    calculated = [{"action": "delete", "table": "T", "name": name,
                   "item_type": "Calculated Column"} for name in ("Multi", "Fenced", "Inline")]
    policy = evaluate_deletion_policy(model, [report], [base])
    assert not policy["ok"]
    assert {v["message"] for v in policy["violations"] if v["rule_id"] == "SMC-D006"} == {
        f"T[Base] is required by retained T[{name}]." for name in ("Multi", "Fenced", "Inline")}
    for action in calculated:
        policy = evaluate_deletion_policy(model, [report], [action])
        assert policy["ok"], policy["errors"]
    policy = evaluate_deletion_policy(model, [report], [base, *calculated])
    assert policy["ok"], policy["errors"]


def test_perspective_members_come_from_declarations_only():
    members = extract_perspective_members(
        "perspective Exec\n\t// perspectiveMeasure Ghost\n\tperspectiveTable Sales\n\t\tperspectiveMeasure Revenue\n")
    assert [(m.perspective, m.table, m.name, m.kind) for m in members] == [
        ("Exec", "Sales", "", "table"), ("Exec", "Sales", "Revenue", "measure")]


def test_kpi_selector_table_creates_no_kpi_limitation(results, payload):
    limitations = results["analysis_limitations"]
    assert not any("KPI Selector" in limitation["source_file"] for limitation in limitations)
    kpi_areas = [limitation for limitation in limitations if limitation["area"] == "KPI expressions"]
    assert {limitation["owner"] for limitation in kpi_areas} == {"measure 'Targets'[Target]"}
    assert {limitation["line"] for limitation in kpi_areas} == {4, 5, 9}
    for name in ("KPI Name", "KPI Flag", "KPI Label", "KPI Count"):
        triggers = row(results, "KPI Selector", name)["review_triggers"]
        assert not any("KPI" in trigger and "Referenced by" in trigger for trigger in triggers), triggers
        assert not any("Unsupported Metadata:" in trigger for trigger in triggers)
    browser_triggers = [trigger for i in payload["items"] for trigger in i["reviewTriggers"]]
    assert not any("KPI Selector.tmdl" in trigger for trigger in browser_triggers)


def test_genuine_kpi_reference_is_evidence_on_the_referenced_item(results, payload):
    goal = row(results, "Sales", "Revenue Goal")
    assert goal["status"] == "NOT USED"
    assert goal["removal_risk"] == "Review"
    # The KPI target is evidence, not a DAX dependency of the owning measure (#92):
    # the item stays Unused with exactly one targeted KPI trigger.
    kpi_triggers = [trigger for trigger in goal["review_triggers"] if "KPI" in trigger]
    assert len(kpi_triggers) == 1
    assert kpi_triggers[0].startswith("Referenced by the KPI target expression of measure 'Targets'[Target] "
                                      "(definition/tables/Targets.tmdl:4)")
    browser = item(payload, "Sales", "Revenue Goal")
    assert browser["usageState"] == "Unused"
    assert browser["deleteSafety"] == "Review"
    # Sales[Revenue] is used directly; the KPI reference must not downgrade it.
    assert row(results, "Sales", "Revenue")["status"] == "USED"


def test_unresolved_dynamic_coverage_still_prevents_safe(results):
    assert not results["coverage"]["complete"]
    shared = results["coverage"]["limitations"]
    assert {limitation["owner"] for limitation in shared} == {
        "calculation item 'Current' in 'Time Intelligence'",
        "calculation item 'YTD' in 'Time Intelligence'",
        "calculation item 'PY' in 'Time Intelligence'",
    }
    unused = [r for r in results["items"] if r["status"] == "NOT USED"]
    assert unused and all(r["removal_risk"] != "Safe" for r in unused)
    assert all(sum(t.startswith("No use found in the selected scope") for t in r["review_triggers"]) == 1
               for r in unused)


# ── #83: calculation-group and perspective dependencies ──────────────────────

def test_selector_shows_parent_table_consumers_separately(results, payload):
    selector = row(results, "Time Intelligence", "Name")
    assert selector["status"] == "NOT USED"
    assert selector["removal_risk"] == "Review"
    assert selector["model_role"] == "Calculation group selector"
    assert selector["table_kind"] == "Calculation group"
    assert selector["table_dependents"] == CONSUMERS
    triggers = selector["review_triggers"]
    assert any(t.startswith("Selector column of calculation group 'Time Intelligence' (3 calculation items: "
                            "Current, YTD, PY)") for t in triggers)
    parent = next(t for t in triggers if t.startswith("Parent table 'Time Intelligence' is required by 10 retained DAX consumers"))
    assert "parent-table references, not direct consumers of this item" in parent
    browser = item(payload, "Time Intelligence", "Name")
    assert browser["modelRole"] == "Calculation group selector"
    assert browser["tableDependentCount"] == 10
    assert browser["tableDependentItems"] == CONSUMERS
    assert browser["measureDependentCount"] == 0  # not pretended to be direct column consumers
    assert browser["usageCount"] == 0  # not pretended to be Report References
    assert browser["usageState"] == "Unused"
    assert browser["deleteSafety"] == "Review"


def test_calculation_group_table_is_identified_and_blocked(results, payload):
    summary = next(t for t in results["table_summaries"] if t["name"] == "Time Intelligence")
    assert summary["table_kind"] == "Calculation group"
    assert summary["role_label"] == "calculation-group"
    assert "No relationships" not in summary["role_reason"]
    assert summary["calculation_items"] == ["Current", "YTD", "PY"]
    assert summary["external_dax_dependents"] == CONSUMERS
    assert summary["cleanup_recommendation"] == "Blocked"
    assert "10 retained DAX consumers" in summary["cleanup_reason"]
    assert "Sales[Revenue Ignoring TI 01]" in summary["cleanup_reason"]
    assert summary["perspectives"] == ["Executive"]
    table = next(t for t in payload["tables"] if t["name"] == "Time Intelligence")
    assert table["tableKind"] == "Calculation group"
    assert table["cleanupRecommendation"] == "Blocked"
    assert table["usageState"] == "Indirect"
    assert table["usageStatus"].startswith("INDIRECT (via: ")
    assert "10 DAX consumers outside this table" in table["usageStatus"]
    assert table["relationshipCount"] == 0
    assert table["externalDaxDependents"] == CONSUMERS


def test_perspective_membership_is_concrete_evidence_not_runtime_use(results, payload):
    goal = row(results, "Sales", "Revenue Goal")
    assert goal["perspectives"] == [
        {"perspective": "Executive", "source_file": "definition/perspectives/Executive.tmdl", "line": 4}]
    assert not any("perspective" in t.lower() for t in goal["review_triggers"])
    assert goal["status"] == "NOT USED"
    assert not any(limitation["area"] == "Perspectives" for limitation in results["analysis_limitations"])
    browser = item(payload, "Sales", "Revenue Goal")
    assert browser["perspectiveMemberships"] == [
        {"perspective": "Executive", "sourceFile": "definition/perspectives/Executive.tmdl"}]
    assert browser["usageCount"] == 0  # membership never counts as a Report Reference
    assert browser["usageState"] == "Unused"
    assert browser["deleteSafety"] == "Review"


def test_deletion_guards_block_whole_group_and_allow_coordinated_final_state(tmp_path):
    policy = evaluate_deletion_policy(MODEL, [REPORT], [
        {"action": "delete", "table": "Time Intelligence", "name": "", "item_type": "table"}])
    assert not policy["ok"]
    table_guards = [v for v in policy["violations"] if v["rule_id"] == "SMC-D006"]
    assert len(table_guards) == 10
    assert {v["message"] for v in table_guards} == {
        f"Table Time Intelligence is required by retained {consumer}." for consumer in CONSUMERS}
    assert any(v["rule_id"] == "SMC-D002" for v in policy["violations"])  # dynamic coverage stays conservative
    snapshot = {p: p.read_bytes() for p in FIXTURE.rglob("*") if p.is_file()}
    assert {p: p.read_bytes() for p in FIXTURE.rglob("*") if p.is_file()} == snapshot

    # Without dynamic limitations, the same parent-table guard is evaluated against
    # the proposed final state: removing the consumers together is allowed.
    model = tmp_path / "Plain.SemanticModel"
    report = tmp_path / "Exec.Report"
    (model / "definition" / "tables").mkdir(parents=True)
    (report / "definition").mkdir(parents=True)
    (model / "definition/tables/Dim.tmdl").write_text("table Dim\n\tcolumn Key\n\t\tdataType: string\n", encoding="utf-8")
    (model / "definition/tables/Sales.tmdl").write_text(
        "table Sales\n\tmeasure Rows = COUNTROWS(Sales)\n\tmeasure Outside = CALCULATE([Rows], ALL('Dim'))\n",
        encoding="utf-8")
    (report / "definition.pbir").write_text(json.dumps({"datasetReference": {"byPath": {"path": "../Plain.SemanticModel"}}}), encoding="utf-8")
    (report / "definition/report.json").write_text("{}", encoding="utf-8")
    whole_table = [{"action": "delete", "table": "Dim", "name": "", "item_type": "table"}]
    blocked = evaluate_deletion_policy(model, [report], whole_table)
    assert [v["message"] for v in blocked["violations"]] == ["Table Dim is required by retained Sales[Outside]."]
    coordinated = whole_table + [{"action": "delete", "table": "Sales", "name": "Outside", "item_type": "Measure"}]
    assert evaluate_deletion_policy(model, [report], coordinated)["ok"]


# ── #91: coordinated whole-group deletion against the simulated final state ──

GROUP_TABLE = [{"action": "delete", "table": "Time Intelligence", "name": "", "item_type": "table"}]
CONSUMER_DELETES = [{"action": "delete", "table": "Sales", "name": consumer[len("Sales["):-1], "item_type": "Measure"}
                    for consumer in CONSUMERS]
GROUP_OWNERS = {
    "calculation item 'Current' in 'Time Intelligence'",
    "calculation item 'YTD' in 'Time Intelligence'",
    "calculation item 'PY' in 'Time Intelligence'",
}
CURRENCY_GROUP = (
    "table 'Currency Conversion'\n"
    "\tcalculationGroup\n"
    "\t\tprecedence: 20\n\n"
    "\t\tcalculationItem Local = SELECTEDMEASURE()\n\n"
    "\tcolumn Currency\n"
    "\t\tdataType: string\n"
    "\t\tsourceColumn: Name\n\n"
    "\tpartition 'Currency Conversion' = calculationGroup\n"
    "\t\tmode: import\n"
    "\t\tsource = calculationGroup\n"
)


def standalone_group_copy(tmp_path, *, other_group=False, keep_hidden=False):
    """Copy the fixture without the item-specific facts that keep a delete at Review.

    The shipped fixture deliberately uses Revenue Ignoring TI 01 in a visual, keeps
    the selector in a perspective and hides Ordinal; those are separate guards
    (SMC-D004/SMC-D005) unrelated to coverage, so the copy removes them.
    ``keep_hidden`` keeps the always-hidden Ordinal column of a real group (#103).
    """
    import shutil

    root = tmp_path / "workspace"
    shutil.copytree(FIXTURE, root)
    model = root / "Models" / "Synthetic Dependencies.SemanticModel"
    report = root / "Reports" / "Executive.Report"
    visual = report / "definition/pages/Overview/visuals/RevenueCard/visual.json"
    data = json.loads(visual.read_text(encoding="utf-8"))
    values = data["visual"]["query"]["queryState"]["Values"]
    values["projections"] = [p for p in values["projections"] if p["queryRef"] == "Sales.Revenue"]
    visual.write_text(json.dumps(data, indent=2), encoding="utf-8")
    perspective = model / "definition/perspectives/Executive.tmdl"
    text = perspective.read_text(encoding="utf-8")
    perspective.write_text(text.split("\tperspectiveTable 'Time Intelligence'")[0].rstrip() + "\n", encoding="utf-8")
    if not keep_hidden:
        group = model / "definition/tables/Time Intelligence.tmdl"
        group.write_text(group.read_text(encoding="utf-8").replace("\t\tisHidden\n", ""), encoding="utf-8")
    if other_group:
        (model / "definition/tables/Currency Conversion.tmdl").write_text(CURRENCY_GROUP, encoding="utf-8")
        definition = model / "definition/model.tmdl"
        definition.write_text(definition.read_text(encoding="utf-8") + "ref table 'Currency Conversion'\n",
                              encoding="utf-8")
    return root, model, report


def test_coordinated_group_deletion_clears_the_groups_own_coverage_gap():
    policy = evaluate_deletion_policy(MODEL, [REPORT], GROUP_TABLE + CONSUMER_DELETES)
    rules = {v["rule_id"] for v in policy["violations"]}
    # In the final state the group file is gone and no retained item references it.
    assert "SMC-D002" not in rules
    assert "SMC-D006" not in rules
    assert policy["scope"]["complete"] is True
    assert {l["owner"] for l in policy["scope"]["cleared_limitations"]} == GROUP_OWNERS
    assert len(policy["scope"]["cleared_limitations"]) == 4
    assert all(l["cleared_reason"] == "Coverage gap owned by 'Time Intelligence' is cleared because the plan "
               "removes the calculation group and every retained parent-table consumer."
               for l in policy["scope"]["cleared_limitations"])
    # Item-specific guards are unchanged: a used consumer still requires review. The
    # cleared gap, the group-structure reasons, perspective membership and the hidden
    # flag of the group's own Ordinal column no longer appear.
    remaining = {(v["rule_id"], v["table"], v["name"]) for v in policy["violations"]}
    assert remaining == {("SMC-D004", "Sales", "Revenue Ignoring TI 01")}


def test_coordinated_group_deletion_is_allowed_when_no_item_guard_remains(tmp_path):
    _, model, report = standalone_group_copy(tmp_path)
    before = {p: p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    policy = evaluate_deletion_policy(model, [report], GROUP_TABLE + CONSUMER_DELETES)
    assert policy["ok"], policy["errors"]
    assert policy["scope"]["complete"] is True
    assert {l["owner"] for l in policy["scope"]["cleared_limitations"]} == GROUP_OWNERS
    # Deleting only the group stays Blocked by the parent-table guard, and its own
    # coverage gap still counts because retained consumers keep referencing it.
    blocked = evaluate_deletion_policy(model, [report], GROUP_TABLE)
    assert not blocked["ok"]
    assert len([v for v in blocked["violations"] if v["rule_id"] == "SMC-D006"]) == 10
    assert any(v["rule_id"] == "SMC-D002" for v in blocked["violations"])
    assert blocked["scope"]["cleared_limitations"] == []
    # Deleting only the consumers keeps the group and its gap.
    consumers_only = evaluate_deletion_policy(model, [report], CONSUMER_DELETES)
    assert any(v["rule_id"] == "SMC-D002" for v in consumers_only["violations"])
    assert {p: p.read_bytes() for p in before} == before


def test_reviewed_plan_preview_explains_cleared_group_coverage(tmp_path):
    from semantic_model_cleaner import change_plan

    _, model, report = standalone_group_copy(tmp_path)
    columns = [{"action": "delete", "table": "Time Intelligence", "name": name, "item_type": "Column"}
               for name in ("Name", "Ordinal")]
    plan = change_plan.create_plan(model, [report], [{"kind": "actions", "actions": columns + CONSUMER_DELETES}])
    assert any(change["path"].endswith("tables/Time Intelligence.tmdl") and change["change"] == "deleted"
               for change in plan["changes"])
    assert plan["coverage"]["final_state"]["complete"] is True
    assert ("Coverage gap owned by 'Time Intelligence' is cleared because the plan removes the calculation "
            "group and every retained parent-table consumer.") in plan["validation"]["limitations"]


def test_retained_other_group_still_blocks_coordinated_deletion(tmp_path):
    _, model, report = standalone_group_copy(tmp_path, other_group=True)
    policy = evaluate_deletion_policy(model, [report], GROUP_TABLE + CONSUMER_DELETES)
    assert not policy["ok"]
    coverage = [v["message"] for v in policy["violations"] if v["rule_id"] == "SMC-D002"]
    assert coverage == ["Dependency checking is incomplete for the calculation item expression of calculation item "
                        "'Local' in 'Currency Conversion' (definition/tables/Currency Conversion.tmdl:5)."]
    assert policy["scope"]["complete"] is False
    assert {l["owner"] for l in policy["scope"]["cleared_limitations"]} == GROUP_OWNERS
    # The retained group's gap keeps the consumers at Review.
    assert {v["name"] for v in policy["violations"] if v["rule_id"] == "SMC-D005"} >= {
        consumer[len("Sales["):-1] for consumer in CONSUMERS}


# ── #103: whole tool-entity deletion, Review by entity type ──────────────────

TI_CONFIRMATION = ("Entity type requires confirmation: 'Time Intelligence' is a calculation group, which this "
                   "tool always asks you to confirm before deletion. Nothing that remains after this plan in the "
                   "model or the selected reports uses it.")


def synthetic_project(tmp_path, tables: dict[str, str]):
    model = tmp_path / "M.SemanticModel"
    report = tmp_path / "R.Report"
    (model / "definition/tables").mkdir(parents=True)
    for name, text in tables.items():
        (model / "definition/tables" / f"{name}.tmdl").write_text(text, encoding="utf-8")
    (report / "definition").mkdir(parents=True)
    (report / "definition/report.json").write_text("{}", encoding="utf-8")
    (report / "definition.pbir").write_text(
        json.dumps({"datasetReference": {"byPath": {"path": "../M.SemanticModel"}}}), encoding="utf-8")
    return model, report


SCENARIO_GROUP = (
    "table Scenario\n"
    "\tcalculationGroup\n"
    "\t\tprecedence: 5\n\n"
    "\t\tcalculationItem Actual = SELECTEDMEASURE()\n\n"
    "\tcolumn Scenario\n"
    "\t\tdataType: string\n"
    "\t\tsourceColumn: Name\n"
    "\t\tsortByColumn: Ordinal\n\n"
    "\tcolumn Ordinal\n"
    "\t\tdataType: int64\n"
    "\t\tisHidden\n"
    "\t\tsourceColumn: Ordinal\n\n"
    "\tpartition Scenario = calculationGroup\n"
    "\t\tmode: import\n"
    "\t\tsource = calculationGroup\n"
)
METRIC_PARAMETER = (
    "table Parameter\n"
    "\tcolumn Parameter\n"
    "\t\tdataType: string\n"
    "\t\tsourceColumn: [Value1]\n"
    "\t\tsortByColumn: 'Parameter Order'\n\n"
    "\tcolumn 'Parameter Fields'\n"
    "\t\tdataType: string\n"
    "\t\tisHidden\n"
    "\t\tsourceColumn: [Value2]\n\n"
    "\tcolumn 'Parameter Order'\n"
    "\t\tdataType: int64\n"
    "\t\tisHidden\n"
    "\t\tsourceColumn: [Value3]\n\n"
    "\tpartition Parameter = calculated\n"
    "\t\tsource =\n"
    "\t\t\t{ (\"Target\", NAMEOF('Sales'[Target]), 0) }\n"
)
SALES = "table Sales\n\tmeasure Existing = 1\n\tmeasure Target = 2\n"


def column_deletes(table, names):
    return [{"action": "delete", "table": table, "name": name, "item_type": "Column"} for name in names]


def assert_no_hidden_reason(policy):
    assert not any("hidden" in v["message"].casefold() for v in policy["violations"])
    assert not any("hidden" in c["message"].casefold() for c in policy["confirmations"])


def test_whole_group_plan_gives_one_entity_type_confirmation_and_no_hidden_trigger(tmp_path):
    # #91 fixture variant without perspective membership (#101) and without the
    # used consumer, keeping the always-hidden Ordinal column of a real group.
    _, model, report = standalone_group_copy(tmp_path, keep_hidden=True)
    policy = evaluate_deletion_policy(model, [report], GROUP_TABLE + CONSUMER_DELETES)
    assert policy["ok"], policy["errors"]
    assert policy["violations"] == []
    assert policy["confirmations"] == [{"kind": "entity_type", "table": "Time Intelligence",
                                        "message": TI_CONFIRMATION}]
    assert_no_hidden_reason(policy)
    assert "used" not in TI_CONFIRMATION.casefold()
    # Selecting every group column instead of the table action is the same plan.
    columns = column_deletes("Time Intelligence", ["Name", "Ordinal"])
    assert evaluate_deletion_policy(model, [report], columns + CONSUMER_DELETES)["confirmations"] == \
        policy["confirmations"]


def test_whole_group_plan_preview_asks_for_entity_type_confirmation(tmp_path):
    from semantic_model_cleaner import change_plan

    _, model, report = standalone_group_copy(tmp_path, keep_hidden=True)
    columns = column_deletes("Time Intelligence", ["Name", "Ordinal"])
    plan = change_plan.create_plan(model, [report], [{"kind": "actions", "actions": columns + CONSUMER_DELETES}])
    assert plan["validation"]["confirmations"] == [TI_CONFIRMATION]
    assert any(change["path"].endswith("tables/Time Intelligence.tmdl") and change["change"] == "deleted"
               for change in plan["changes"])


def test_partial_group_plan_keeps_item_alone_review_reasons(tmp_path):
    _, model, report = standalone_group_copy(tmp_path, keep_hidden=True)
    analysis = analyzer.analyze(model.parent, model_paths=[model], report_paths=[report])
    ordinal = row(analysis, "Time Intelligence", "Ordinal")
    # Retained consumers exist, so the analysis view keeps today's reasons.
    assert ordinal["review_triggers"] == ordinal["standalone_review_triggers"]
    assert ordinal["review_basis"] == "evidence"
    for actions in (column_deletes("Time Intelligence", ["Ordinal"]),
                    column_deletes("Time Intelligence", ["Ordinal"]) + CONSUMER_DELETES):
        policy = evaluate_deletion_policy(model, [report], actions)
        assert not policy["ok"]
        assert policy["confirmations"] == []
        messages = {v["name"]: v["message"] for v in policy["violations"] if v["rule_id"] == "SMC-D005"}
        assert messages["Ordinal"] == ("Review required for Time Intelligence[Ordinal]: "
                                       + "; ".join(ordinal["review_triggers"]))
        assert "Item is hidden" in messages["Ordinal"]
        assert "Delete the whole group through a reviewed table plan" in messages["Ordinal"]


def test_unused_calculation_group_is_review_by_entity_type(tmp_path):
    model, report = synthetic_project(tmp_path, {"Scenario": SCENARIO_GROUP, "Sales": SALES})
    analysis = analyzer.analyze(tmp_path, model_paths=[model], report_paths=[report])
    trigger = ("Entity type requires confirmation: 'Scenario' is a calculation group, which this tool always asks "
               "you to confirm before deletion. Nothing in the model or the selected reports uses it.")
    for name in ("Scenario", "Ordinal"):
        group_row = row(analysis, "Scenario", name)
        assert group_row["removal_risk"] == "Review"
        assert group_row["review_triggers"] == [trigger]
        assert group_row["review_basis"] == "entity_type"
    assert "Item is hidden" in row(analysis, "Scenario", "Ordinal")["standalone_review_triggers"]
    table = next(t for t in analysis["table_summaries"] if t["name"] == "Scenario")
    assert (table["cleanup_recommendation"], table["cleanup_reason"], table["review_basis"]) == (
        "Review", trigger, "entity_type")
    assert analysis["summary"]["review_entity_type"] == 2
    browser = webapp._serialize_results(analysis, model_paths=[str(model)])
    ordinal = item(browser, "Scenario", "Ordinal")
    assert (ordinal["reviewBasis"], ordinal["entityKind"], ordinal["reviewTriggers"]) == (
        "entity_type", "calculation group", [trigger])
    assert next(t for t in browser["tables"] if t["name"] == "Scenario")["cleanupReason"] == trigger

    whole = evaluate_deletion_policy(model, [report], column_deletes("Scenario", ["Scenario", "Ordinal"]))
    assert whole["ok"], whole["errors"]
    assert [c["table"] for c in whole["confirmations"]] == ["Scenario"]
    assert_no_hidden_reason(whole)
    partial = evaluate_deletion_policy(model, [report], column_deletes("Scenario", ["Ordinal"]))
    assert not partial["ok"] and partial["confirmations"] == []
    assert any(v["rule_id"] == "SMC-D005" and v["name"] == "Ordinal" and "Item is hidden" in v["message"]
               for v in partial["violations"])


def test_unused_field_parameter_is_review_by_entity_type(tmp_path):
    from semantic_model_cleaner import change_plan

    model, report = synthetic_project(tmp_path, {"Parameter": METRIC_PARAMETER, "Sales": SALES})
    analysis = analyzer.analyze(tmp_path, model_paths=[model], report_paths=[report])
    trigger = ("Entity type requires confirmation: 'Parameter' is a field parameter, which this tool always asks "
               "you to confirm before deletion. Nothing in the model or the selected reports uses it.")
    names = ["Parameter", "Parameter Fields", "Parameter Order"]
    for name in names:
        parameter_row = row(analysis, "Parameter", name)
        assert (parameter_row["removal_risk"], parameter_row["review_triggers"]) == ("Review", [trigger])
    # Guards outside the entity are unchanged: the NAMEOF target stays Caution.
    assert row(analysis, "Sales", "Target")["removal_risk"] == "Caution"
    assert row(analysis, "Sales", "Existing")["removal_risk"] == "Safe"
    assert analysis["summary"]["review_entity_type"] == 3
    assert analysis["summary"]["review_evidence"] == 0

    whole = evaluate_deletion_policy(model, [report], column_deletes("Parameter", names))
    assert whole["ok"], whole["errors"]
    assert whole["confirmations"] == [{"kind": "entity_type", "table": "Parameter", "message": trigger.replace(
        "Nothing in the model", "Nothing that remains after this plan in the model")}]
    assert_no_hidden_reason(whole)
    plan = change_plan.create_plan(model, [report], [{"kind": "actions",
                                                     "actions": column_deletes("Parameter", names)}])
    assert plan["validation"]["confirmations"] == [whole["confirmations"][0]["message"]]
    # Partial plans keep today's item-alone behaviour.
    hidden_only = evaluate_deletion_policy(model, [report], column_deletes("Parameter", ["Parameter Fields"]))
    assert [v["message"] for v in hidden_only["violations"]] == [
        "Review required for Parameter[Parameter Fields]: Item is hidden"]
    assert hidden_only["confirmations"] == []
    assert evaluate_deletion_policy(model, [report], column_deletes("Parameter", ["Parameter"]))["ok"]
    # Deleting the NAMEOF target together with the parameter stays guarded.
    target = [{"action": "delete", "table": "Sales", "name": "Target", "item_type": "Measure"}]
    assert any(v["rule_id"] == "SMC-D006" for v in evaluate_deletion_policy(
        model, [report], column_deletes("Parameter", names) + target)["violations"])


def test_used_field_parameter_is_not_review_by_entity_type(tmp_path):
    model, report = synthetic_project(tmp_path, {"Parameter": METRIC_PARAMETER, "Sales": SALES})
    analysis = analyzer.analyze(tmp_path, model_paths=[model], report_paths=[report])
    assert analysis["summary"]["review_entity_type"] == 3
    # A retained measure that references the parameter table keeps evidence-based reasons.
    (model / "definition/tables/Sales.tmdl").write_text(
        SALES + "\tmeasure Picked = SELECTEDVALUE(Parameter[Parameter])\n", encoding="utf-8")
    analysis = analyzer.analyze(tmp_path, model_paths=[model], report_paths=[report])
    assert analysis["summary"]["review_entity_type"] == 0
    fields = row(analysis, "Parameter", "Parameter Fields")
    assert fields["review_triggers"] == ["Item is hidden"]
    assert fields["review_basis"] == "evidence"


# ── #84: analysis limitations are a separate surface ─────────────────────────

def test_limitation_counts_use_distinct_units(results, payload):
    summary = results["analysis_limitation_summary"]
    limitations = payload["analysisLimitations"]
    # 3 calculation items + 1 calculation-item format string (shared, dynamic);
    # KPI target, status and trend expressions (targeted).
    assert summary["distinct_count"] == limitations["distinctCount"] == 7
    assert summary["shared_count"] == limitations["sharedCount"] == 4
    assert summary["targeted_count"] == limitations["targetedCount"] == 3
    unused = [r for r in results["items"] if r["status"] == "NOT USED"]
    assert summary["affected_item_count"] == limitations["affectedItemCount"] == len(unused)
    # Repeating a shared reason on every item must not inflate the distinct count.
    assert sum(len(r["analysis_limitation_ids"]) for r in results["items"]) > summary["distinct_count"]
    assert limitations["coverageComplete"] is False


def test_each_limitation_explains_feature_location_and_effect(payload):
    by_owner = {(l["area"], l["owner"], l["line"]): l for l in payload["analysisLimitations"]["limitations"]}
    format_string = by_owner[("Format string definitions", "calculation item 'YTD' in 'Time Intelligence'", 8)]
    assert format_string["feature"] == "calculation item format-string expression"
    assert format_string["message"] == (
        "Dependency checking is incomplete for the calculation item format-string expression of "
        "calculation item 'YTD' in 'Time Intelligence' (definition/tables/Time Intelligence.tmdl:8).")
    assert "invalid" not in format_string["message"].lower()
    assert format_string["scope"] == "shared"
    assert format_string["checked"] and format_string["unchecked"] and format_string["effect"]
    kpi = by_owner[("KPI expressions", "measure 'Targets'[Target]", 4)]
    assert kpi["scope"] == "targeted"
    assert kpi["targets"] == ["Sales[Revenue Goal]"]
    assert kpi["affectedItems"] == [{"type": "Measure", "table": "Sales", "name": "Revenue Goal"}]
    assert kpi["effect"] == "Only the referenced items are affected; they stay at Review."
    trend = by_owner[("KPI expressions", "measure 'Targets'[Target]", 9)]
    assert trend["scope"] == "targeted" and trend["targets"] == [] and trend["affectedItemCount"] == 0


def test_report_health_excludes_model_limitations(payload):
    keys = {group["key"] for group in payload["reportHealth"]["groups"]}
    assert "unsupported_metadata" not in keys
    assert payload["reportHealth"]["totalIssueCount"] == 0
    assert payload["reportHealth"]["signalCounts"] == {"staleReferences": 0, "brokenModelReferences": 0}


def test_item_summaries_lead_with_concrete_reason_and_collapse_shared_gaps(results):
    selector = row(results, "Time Intelligence", "Name")
    triggers = selector["review_triggers"]
    assert triggers[0].startswith("Selector column of calculation group")
    assert triggers[-1].startswith("No use found in the selected scope, but 4 analysis limitations (")
    assert "+2 more" in triggers[-1]
    assert selector["analysis_limitation_ids"] == [l["id"] for l in results["coverage"]["limitations"]]


def test_exports_and_ci_agree_with_browser_counts(results, payload, tmp_path):
    exported = json.loads(analyzer.format_json_output(results))
    assert exported["analysisLimitationSummary"]["distinct_count"] == payload["analysisLimitations"]["distinctCount"]
    assert [l["id"] for l in exported["analysisLimitations"]] == [l["id"] for l in payload["analysisLimitations"]["limitations"]]
    selector = next(i for i in exported["items"] if i["table"] == "Time Intelligence" and i["name"] == "Name")
    assert selector["modelRole"] == "Calculation group selector"
    assert selector["tableDependents"] == CONSUMERS
    assert selector["analysisLimitationIds"]
    markdown = analyzer.format_full(results)
    assert "## Analysis Limitations" in markdown
    assert "7 distinct limitations" in markdown
    assert "KPI Selector" not in markdown.split("## Summary")[0]
    code, check = ci.run_check(FIXTURE, model_path=MODEL, report_paths=[REPORT])
    assert code == 1
    smc003 = [f for f in check["findings"] if f["rule_id"] == "SMC003"]
    assert len(smc003) == 7
    assert not any("KPI Selector" in f["path"] for f in smc003)
    assert any(f["message"].startswith("Analysis limitation (KPI expressions): dependency checking is incomplete for the "
                                       "KPI target expression of measure 'Targets'[Target]") for f in smc003)
